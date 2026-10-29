"""Rapports de simulation : feuille Excel (.xlsx) et document PDF.

Les chiffres viennent du résultat natif (`RunResult`) ; ce module ne calcule que les
taux d'affichage. Si le résultat porte un profil énergie (`run_scenario(..., energy_profile=...)`),
l'énergie par message et le courant moyen figurent dans le rapport.
Les textes issus du scénario (nom, groupes) sont neutralisés contre l'injection de formules
tableur (`=`, `+`, `-`, `@` en tête de cellule).
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Union

from lsim import __version__
from lsim.energy_report import profile_name
from lsim.export import DeviceRow, RunResult

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

_ENERGY_SUMMARY_LABEL = "Profil énergie (classe A, un message par période)"


def _safe_text(value: str) -> str:
    """Neutralise une chaîne pouvant être interprétée comme formule tableur."""
    if value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _summary_rows(result: RunResult) -> list[tuple[str, object]]:
    t = result.totals
    g = result.gateway
    pdr = f"{result.pdr * 100:.2f} %" if result.pdr is not None else "n/a"
    rows: list[tuple[str, object]] = [
        ("Scénario", _safe_text(result.scenario_name)),
        ("Empreinte SHA-256 du scénario", result.scenario_sha256),
        ("Graine", result.seed),
        ("Version lsim (Python)", __version__),
        ("Émises", t.sent),
        ("Reçues", t.received),
        ("PDR", pdr),
        ("Pertes par sensibilité", t.lost_sensitivity),
        ("Pertes par collision", t.lost_collision),
        ("Temps d'antenne total (ns)", t.airtime_ns),
    ]
    if g is not None:
        rows += [
            ("Charge passerelle : durée occupée (ns)", g.busy_ns),
            ("Charge passerelle : pic de trames simultanées", g.peak_concurrent),
            ("Charge passerelle : taux d'occupation", round(g.utilization, 6)),
        ]
    for name, g in zip(result.gateway_names, result.gateway_loads):
        rows += [
            (f"Passerelle {name} : trames audibles", g.audible_frames),
            (f"Passerelle {name} : taux d'occupation", round(g.utilization, 6)),
            (f"Passerelle {name} : pic de trames simultanées", g.peak_concurrent),
        ]
    if result.mesh is not None:
        m = result.mesh
        rows += [
            ("Mesh : trames décodées à l'accès", m.frames),
            ("Mesh : trames remises au LNS", m.delivered),
            ("Mesh : remises sans saut mesh", m.direct),
            ("Mesh : remises après saut mesh", m.via_mesh),
            ("Mesh : copies confiées au mesh", m.copies_on_mesh),
            ("Mesh : émissions mesh", m.transmissions),
            ("Mesh : copies supprimées (doublon)", m.suppressed),
            ("Mesh : copies sans route", m.no_route),
            ("Mesh : trames perdues sur le mesh", m.lost_mesh),
            ("Mesh : descente dans RX1", m.dl_rx1),
            ("Mesh : descente dans RX2 seulement", m.dl_rx2),
            ("Mesh : descente hors fenêtres", m.dl_missed),
        ] + [(f"Mesh : sauts perdus ({reason})", n) for reason, n in m.hop_losses]
    if result.energy_profile is not None:
        rows.append((_ENERGY_SUMMARY_LABEL, profile_name(result.energy_profile)))
    return rows


def _fmt_mj(value: Optional[float]) -> object:
    return round(value * 1e3, 6) if value is not None else "n/a"


def _energy_cells(r: DeviceRow) -> tuple[object, object]:
    """(mJ par message, courant moyen µA) d'une ligne ; « n/a » si non calculé."""
    mj = _fmt_mj(r.energy_per_message_j)
    ua = round(r.average_current_ua, 6) if r.average_current_ua is not None else "n/a"
    return mj, ua


def write_xlsx(result: RunResult, path: Union[str, Path]) -> None:
    """Écrit un classeur : « Synthèse » et « Dispositifs » (une ligne par dispositif)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    ws.append(["Indicateur", "Valeur"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for label, value in _summary_rows(result):
        ws.append([label, value])
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 66

    dv = wb.create_sheet("Dispositifs")
    headers = ["Index", "Groupe", "Index dans le groupe", "x (m)", "y (m)", "SF",
               "Émises", "Reçues", "Pertes sensibilité", "Pertes collision", "PDR",
               "Temps d'antenne (ns)"]
    with_energy = result.energy_profile is not None
    if with_energy:
        headers += ["Énergie par message (mJ)", "Courant moyen (µA)"]
    dv.append(headers)
    for c in dv[1]:
        c.font = Font(bold=True)
    for r in result.rows:
        line = [
            r.device_index, _safe_text(r.group), r.index_in_group, r.x_m, r.y_m,
            r.spreading_factor, r.sent, r.received, r.lost_sensitivity, r.lost_collision,
            r.pdr if r.pdr is not None else "", r.airtime_ns,
        ]
        if with_energy:
            line += list(_energy_cells(r))
        dv.append(line)
    dv.freeze_panes = "A2"
    dv.auto_filter.ref = dv.dimensions
    wb.save(path)


def write_pdf(result: RunResult, path: Union[str, Path]) -> None:
    """Écrit un rapport PDF : synthèse et tableau des 200 premiers dispositifs."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    styles = getSampleStyleSheet()
    story = [Paragraph("Rapport de simulation LoRaWAN", styles["Title"]), Spacer(1, 12)]
    summary = [[Paragraph(escape(str(k)), styles["Normal"]),
                Paragraph(escape(str(v)), styles["Normal"])]
               for k, v in _summary_rows(result)]
    t = Table(summary, colWidths=[210, 290])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [t, Spacer(1, 18)]

    shown = result.rows[:200]
    story.append(Paragraph(f"Dispositifs ({len(shown)} sur {len(result.rows)} affichés)",
                           styles["Heading2"]))
    with_energy = result.energy_profile is not None
    rows = [["Index", "Groupe", "SF", "Émises", "Reçues", "PDR"]]
    if with_energy:
        rows[0] += ["mJ/msg", "µA moy."]
    for r in shown:
        line = [str(r.device_index), escape(r.group)[:24], str(r.spreading_factor),
                str(r.sent), str(r.received),
                f"{r.pdr * 100:.1f} %" if r.pdr is not None else "n/a"]
        if with_energy:
            mj, ua = _energy_cells(r)
            line += [f"{mj:.3f}" if isinstance(mj, float) else "n/a",
                     f"{ua:.3f}" if isinstance(ua, float) else "n/a"]
        rows.append(line)
    dt = Table(rows, repeatRows=1)
    dt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                            ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    story.append(dt)
    SimpleDocTemplate(str(path), pagesize=A4).build(story)
